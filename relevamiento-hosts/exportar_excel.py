"""Arma base_hosts_uruguay.xlsx con lo que juntaron los scripts y la investigación de canales.

Se puede correr las veces que haga falta: actualiza las hojas de datos y respeta lo que
se haya escrito a mano en las hojas "Canales" y "Seguimiento".
"""
import os
from collections import defaultdict

from openpyxl.styles import Font

from comun import CARPETA, abrir_excel, cargar_progreso, escribir_hoja

SALIDA = os.path.join(CARPETA, "base_hosts_uruguay.xlsx")
MIN_PROFESIONAL = 3  # anfitriones con 3 o más anuncios: casi siempre administradoras o inversores

COLS_AIRBNB = [
    ("departamento", "Departamento", 13), ("tipo_y_zona", "Tipo y zona", 34),
    ("nombre_anuncio", "Nombre del anuncio", 40), ("titulo", "Título en búsqueda", 30),
    ("subtitulo", "Detalle", 30), ("capacidad", "Capacidad", 28), ("calificacion", "Calificación", 11),
    ("resenas", "Reseñas", 9), ("precio_noche", "Precio/noche", 12), ("favorito_huespedes", "Favorito huéspedes", 10),
    ("host_nombre", "Anfitrión (nombre público)", 24), ("superhost", "Superanfitrión", 10),
    ("host_anios", "Años como anfitrión", 10), ("host_anuncios", "Anuncios del anfitrión (en el relevamiento)", 14),
    ("host_tipo", "Tipo de anfitrión", 16), ("registro", "N° registro", 16),
    ("lat", "Lat", 10), ("lng", "Lng", 10), ("url", "URL anuncio", 40), ("host_url", "URL perfil anfitrión", 40),
    ("error", "Error al leer", 16),
]

COLS_HOSTS = [
    ("host_nombre", "Anfitrión (nombre público)", 26), ("host_tipo", "Tipo", 16),
    ("host_anuncios", "Anuncios", 10), ("departamentos", "Departamentos", 24), ("superhost", "Superanfitrión", 10),
    ("host_anios", "Años como anfitrión", 10), ("resenas_total", "Reseñas (suma)", 12),
    ("host_url", "URL perfil anfitrión", 42), ("ejemplo_url", "URL de un anuncio", 40),
]

COLS_BOOKING = [
    ("departamento", "Departamento", 13), ("localidad_busqueda", "Localidad", 16), ("nombre", "Nombre", 36),
    ("tipo", "Tipo", 16), ("estrellas", "Estrellas", 9), ("direccion", "Dirección", 40),
    ("puntaje_num", "Puntaje", 8), ("resenas", "Reseñas", 9), ("gestionado_por", "Gestionado por", 26),
    ("particular", "Particular", 10), ("profesional", "Profesional", 10),
    ("telefono", "Teléfono (Google)", 16), ("web", "Web (Google)", 30), ("google_maps_url", "Google Maps", 30),
    ("google_estado", "Estado en Google", 14), ("fuera_de_zona", "Fuera de los 3 deptos", 10),
    ("lat", "Lat", 10), ("lng", "Lng", 10), ("url", "URL Booking", 40), ("error", "Error al leer", 16),
]

COLS_CANALES = [
    ("tipo", "Tipo de canal", 22), ("nombre", "Nombre", 34), ("zona", "Zona", 20),
    ("por_que", "Por qué sirve", 50), ("como_contactar", "Cómo contactar / próximo paso", 50),
    ("fuente", "Fuente", 50), ("estado", "Estado", 14), ("notas", "Notas", 30),
]

CANALES = [
    {"tipo": "Comunidad de anfitriones", "nombre": "Grupo de anfitriones (propuesta de regulación de alojamiento temporario)",
     "zona": "Nacional", "por_que": "Más de 100 propietarios organizados que ya reclaman por condiciones del sector; son voceros naturales.",
     "como_contactar": "Ubicar a los firmantes/voceros de la propuesta (salieron en prensa) y proponer una reunión.",
     "fuente": "https://enperspectiva.uy/wp-content/uploads/2024/05/Grupo-de-anfitriones.-Propuesta-Regulacion-Alojamiento-Temporario.pdf"},
    {"tipo": "Comunidad de anfitriones", "nombre": "Club de anfitriones de Airbnb de Montevideo (grupo de Facebook)",
     "zona": "Montevideo", "por_que": "Más de 100 hosts de Montevideo en un grupo privado.",
     "como_contactar": "Pedir ingreso al grupo con perfil real; participar antes de presentar el servicio. Respetar las reglas del grupo.",
     "fuente": "https://news.airbnb.com/latin-american-hosts-launch-8-home-sharing-clubs/"},
    {"tipo": "Comunidad de anfitriones", "nombre": "Grupos de Facebook/WhatsApp de alquileres por balneario",
     "zona": "Rocha y Maldonado", "por_que": "Cada balneario (Punta del Diablo, La Paloma, La Pedrera, Piriápolis...) tiene grupos de alquileres de temporada donde publican los propios dueños.",
     "como_contactar": "Buscar en Facebook \"alquileres Punta del Diablo\", \"alquileres La Paloma\", etc. y sumarse.", "fuente": ""},
    {"tipo": "Gestor profesional", "nombre": "Hoster", "zona": "Montevideo y Punta del Este",
     "por_que": "Administra propiedades de terceros en Airbnb desde 2019: un acuerdo cubre muchas propiedades.",
     "como_contactar": "Contacto comercial por su web / LinkedIn.", "fuente": "https://uy.linkedin.com/in/vicocampbell"},
    {"tipo": "Inmobiliaria", "nombre": "Santos Dumont Inmobiliaria", "zona": "Punta del Este",
     "por_que": "Hace alquileres temporales y administración de inmuebles.", "como_contactar": "Contacto comercial de su ficha.",
     "fuente": "https://www.zonaprop.com.ar/inmobiliarias/santos-dumont-inmobiliaria_30815144-inmuebles.html"},
    {"tipo": "Inmobiliaria", "nombre": "Crucero Real Estate", "zona": "Punta del Este",
     "por_que": "16 años en alquileres turísticos de temporada.", "como_contactar": "Contacto comercial de su web.",
     "fuente": "https://noticias.perfil.com/noticias/empresas-y-protagonistas/lorena-miraballes-kukurian-de-crucero-real-estate-el-mercado-inmobiliario-de-punta-del-este-con-un-enfoque-amigable-y-personalizado.phtml"},
    {"tipo": "Inmobiliaria", "nombre": "Justo Acá Propiedades e Inversiones", "zona": "Punta del Este",
     "por_que": "Asesora a inversores que compran para renta temporaria.", "como_contactar": "Contacto comercial de su web.",
     "fuente": "https://www.montevideo.com.uy/Especiales/Justo-Aca-amplia-su-presencia-y-llega-a-Punta-del-Este-con-proyectos-de-alto-valor-uc940033"},
    {"tipo": "Inmobiliaria", "nombre": "Antonio Mieres Negocios Inmobiliarios", "zona": "Punta del Este",
     "por_que": "Inmobiliaria histórica (desde 1975) con cartera de alquileres.", "como_contactar": "Contacto comercial de su web.",
     "fuente": "https://uy.linkedin.com/in/mieres"},
    {"tipo": "Inmobiliaria", "nombre": "Punta Ballena Inmobiliaria", "zona": "Punta Ballena",
     "por_que": "Inmobiliaria local de la zona.", "como_contactar": "Contacto comercial de su ficha.",
     "fuente": "https://www.jamesedition.com/offices/real_estate/punta-ballena-inmobiliaria-260506"},
    {"tipo": "Asociación", "nombre": "ADIPE / CIDEM (inmobiliarias y desarrolladores de Punta del Este)",
     "zona": "Maldonado", "por_que": "Agrupa a las inmobiliarias de Punta del Este: un solo contacto llega a muchas.",
     "como_contactar": "Pedir reunión con la directiva para presentar el servicio a los socios.",
     "fuente": "https://www.hosteltur.com/lat/120880_punta-inmobiliarias-acusan-baja-actividad-30-40.html"},
    {"tipo": "Portal", "nombre": "InfoCasas / Mercado Libre (alquiler temporal)", "zona": "Rocha y Maldonado",
     "por_que": "Muchas inmobiliarias de Rocha publican ahí con ficha comercial y teléfono.",
     "como_contactar": "Revisar las fichas de inmobiliarias (no de particulares) y sumarlas a esta hoja.",
     "fuente": "https://inmuebles.mercadolibre.com.ar/alquiler-temporal/uruguay/rocha/"},
    {"tipo": "Organismo", "nombre": "Ministerio de Turismo (Mintur)", "zona": "Nacional",
     "por_que": "Impulsa el registro de viviendas de uso turístico; las ligas de fomento y cámaras locales trabajan con ellos.",
     "como_contactar": "Consultar el registro público de alojamientos y si hay datos abiertos.",
     "fuente": "https://www.elobservador.com.uy/cafe-y-negocios/alquileres-airbnb-y-booking-comision-diputados-aprobo-proyecto-que-busca-regular-la-vivienda-turistica-n5960239"},
    {"tipo": "Asociación", "nombre": "Ligas de fomento y centros comerciales locales", "zona": "Rocha y Maldonado",
     "por_que": "La Paloma, La Pedrera, Punta del Diablo, Piriápolis tienen ligas/centros comerciales que nuclean a los alojamientos.",
     "como_contactar": "Contactar a cada liga y ofrecer una charla para socios.", "fuente": ""},
]

COLS_SEGUIMIENTO = [
    ("contacto", "Contacto (negocio / inmobiliaria)", 32), ("origen", "Origen (Booking, inmobiliaria, formulario...)", 22),
    ("zona", "Zona", 16), ("canal", "Canal usado", 16), ("fecha", "Fecha contacto", 14), ("respuesta", "Respuesta", 20),
    ("proximo", "Próximo paso", 30), ("consentimiento", "¿Dio consentimiento para recibir info? (Ley 18.331)", 18),
    ("notas", "Notas", 40),
]

RESUMEN = [
    ("Relevamiento de alojamientos — Montevideo, Maldonado y Rocha", ""),
    ("", ""),
    ("Tamaño del mercado (datos de terceros, junio 2026)", ""),
    ("Anuncios activos de Airbnb en Uruguay", "20.472"),
    ("Punta del Este", "4.471"),
    ("Montevideo", "2.316"),
    ("Punta del Diablo", "1.186"),
    ("Pinares - Las Delicias", "958"),
    ("Punta Ballena", "754"),
    ("La Paloma", "729"),
    ("Fuente", "https://www.guestfavorites.com/airbnb-occupancy-rates-by-city-in-uruguay"),
    ("", ""),
    ("Comisión de Airbnb", ""),
    ("Antes", "3% al anfitrión (+14-16% al huésped)"),
    ("Ahora", "15,5% todo al anfitrión; migración país por país hasta fin de 2026"),
    ("Fuente", "https://www.hostfully.com/es/blog/tarifas-para-anfitriones-de-airbnb-la-tarifa-del-155-solo-para-anfitriones-explicada-2026/"),
    ("", ""),
    ("Este relevamiento", ""),
]


def armar_hosts(anuncios):
    por_host = defaultdict(list)
    for a in anuncios:
        if a.get("host_id"):
            por_host[a["host_id"]].append(a)
    hosts = []
    for lista in por_host.values():
        n = len(lista)
        tipo = "Profesional (3+)" if n >= MIN_PROFESIONAL else ("2 anuncios" if n == 2 else "Particular")
        for a in lista:
            a["host_anuncios"] = n
            a["host_tipo"] = tipo
        p = lista[0]
        hosts.append({
            "host_nombre": p.get("host_nombre", ""), "host_tipo": tipo, "host_anuncios": n,
            "departamentos": ", ".join(sorted({a.get("departamento", "") for a in lista if a.get("departamento")})),
            "superhost": any(a.get("superhost") for a in lista), "host_anios": p.get("host_anios", ""),
            "resenas_total": sum(a.get("resenas") or 0 for a in lista if isinstance(a.get("resenas"), int)),
            "host_url": p.get("host_url", ""), "ejemplo_url": p.get("url", ""),
        })
    hosts.sort(key=lambda h: -h["host_anuncios"])
    return hosts


def main():
    airbnb = list((cargar_progreso("airbnb").get("anuncios") or {}).values())
    booking = list((cargar_progreso("booking").get("alojamientos") or {}).values())
    for a in airbnb:
        a.setdefault("departamento", a.get("zona_busqueda", ""))
    for b in booking:
        b.setdefault("departamento", b.get("zona_busqueda", ""))
    hosts = armar_hosts(airbnb)
    airbnb.sort(key=lambda a: (a.get("departamento", ""), -(a.get("host_anuncios") or 0)))
    booking.sort(key=lambda b: (b.get("departamento", ""), b.get("localidad_busqueda", "")))

    libro = abrir_excel(SALIDA)

    if "Resumen" in libro.sheetnames:
        del libro["Resumen"]
    hoja = libro.create_sheet("Resumen")
    filas = list(RESUMEN)
    for depto in ("Montevideo", "Maldonado", "Rocha"):
        filas.append((f"Airbnb — {depto}", sum(1 for a in airbnb if a.get("departamento") == depto)))
        filas.append((f"Booking — {depto}", sum(1 for b in booking if b.get("departamento") == depto and not b.get("fuera_de_zona"))))
    filas.append(("Anfitriones de Airbnb identificados", len(hosts)))
    filas.append((f"Anfitriones profesionales ({MIN_PROFESIONAL}+ anuncios)", sum(1 for h in hosts if h["host_tipo"].startswith("Profesional"))))
    filas.append(("Alojamientos de Booking con teléfono comercial", sum(1 for b in booking if b.get("telefono"))))
    for f in filas:
        hoja.append(list(f))
    hoja.column_dimensions["A"].width = 48
    hoja.column_dimensions["B"].width = 70
    for celda in hoja["A"]:
        if celda.value and celda.row in (1, 3, 13, 18):
            celda.font = Font(bold=True, size=14 if celda.row == 1 else 12)

    escribir_hoja(libro, "Airbnb", COLS_AIRBNB, airbnb)
    escribir_hoja(libro, "Airbnb - Anfitriones", COLS_HOSTS, hosts)
    escribir_hoja(libro, "Booking", COLS_BOOKING, booking)
    if "Canales" not in libro.sheetnames:
        escribir_hoja(libro, "Canales", COLS_CANALES, CANALES)
    if "Seguimiento" not in libro.sheetnames:
        escribir_hoja(libro, "Seguimiento", COLS_SEGUIMIENTO, [])

    orden = ["Resumen", "Canales", "Booking", "Airbnb - Anfitriones", "Airbnb", "Seguimiento"]
    libro._sheets = [libro[n] for n in orden] + [h for h in libro._sheets if h.title not in orden]
    libro.active = 0
    libro.save(SALIDA)
    print(f"Excel listo: {SALIDA}")
    print(f"  Airbnb: {len(airbnb)} anuncios, {len(hosts)} anfitriones | Booking: {len(booking)} alojamientos")


if __name__ == "__main__":
    main()
