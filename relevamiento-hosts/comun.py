"""Utilidades compartidas: zonas, guardado de progreso y exportación a Excel."""
import json
import os
import random
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CARPETA = os.path.dirname(os.path.abspath(__file__))
DATOS = os.path.join(CARPETA, "datos")
os.makedirs(DATOS, exist_ok=True)

# Rectángulos aproximados de cada departamento (lat_sur, lng_oeste, lat_norte, lng_este).
# Se usan para recorrer el mapa de Airbnb por cuadrículas.
ZONAS = {
    "Montevideo": (-34.94, -56.43, -34.70, -56.03),
    "Maldonado": (-35.00, -55.45, -34.30, -54.55),
    "Rocha": (-34.75, -54.55, -33.65, -53.35),
}

# Localidades para buscar en Booking (Booking corta en ~1000 resultados por búsqueda,
# así que se busca por localidad y después se eliminan duplicados).
LOCALIDADES_BOOKING = {
    "Montevideo": ["Montevideo"],
    "Maldonado": [
        "Punta del Este", "La Barra", "Manantiales", "José Ignacio", "Punta Ballena",
        "Piriápolis", "Maldonado", "San Carlos", "Pan de Azúcar", "Solís", "Las Flores",
        "Playa Verde", "Portezuelo", "Pinares", "Aiguá", "Garzón",
    ],
    "Rocha": [
        "La Paloma", "La Pedrera", "Punta del Diablo", "Cabo Polonio", "Barra de Valizas",
        "Aguas Dulces", "Punta Rubia", "La Aguada", "Arachania", "Antoniópolis",
        "La Coronilla", "Barra del Chuy", "Chuy", "Rocha", "Castillos", "Lascano",
    ],
}


def pausa(minimo=2.0, maximo=5.0):
    """Espera aleatoria para no saturar los sitios ni disparar bloqueos."""
    time.sleep(random.uniform(minimo, maximo))


def cargar_progreso(nombre):
    ruta = os.path.join(DATOS, f"{nombre}.json")
    if os.path.exists(ruta):
        with open(ruta, encoding="utf-8") as f:
            return json.load(f)
    return {}


def guardar_progreso(nombre, datos):
    ruta = os.path.join(DATOS, f"{nombre}.json")
    tmp = ruta + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=1)
    os.replace(tmp, ruta)


def departamento_por_coordenadas(lat, lng):
    """Devuelve el departamento cuyo rectángulo contiene el punto, o '' si ninguno."""
    if lat is None or lng is None:
        return ""
    for depto, (s, o, n, e) in ZONAS.items():
        if s <= lat <= n and o <= lng <= e:
            return depto
    return ""


ENCABEZADO = PatternFill("solid", fgColor="1F4E78")


def escribir_hoja(libro, titulo, columnas, filas):
    """Reemplaza (o crea) una hoja con encabezados formateados y las filas dadas.

    columnas: lista de (clave, título, ancho)."""
    if titulo in libro.sheetnames:
        del libro[titulo]
    hoja = libro.create_sheet(titulo)
    for i, (_, nombre, ancho) in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=i, value=nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    for fila in filas:
        hoja.append([fila.get(clave, "") for clave, _, _ in columnas])
    hoja.freeze_panes = "A2"
    if filas:
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(filas) + 1}"
    for i, (clave, _, _) in enumerate(columnas, start=1):
        if clave.startswith("url") or clave.endswith("_url"):
            for fila in range(2, len(filas) + 2):
                c = hoja.cell(row=fila, column=i)
                if c.value:
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
    return hoja


def abrir_excel(ruta):
    if os.path.exists(ruta):
        return load_workbook(ruta)
    libro = Workbook()
    del libro[libro.active.title]
    return libro
